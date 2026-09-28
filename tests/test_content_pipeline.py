import pytest
from app.services.content_sanitizer import ContentSanitizer
from app.services.question_validator import QuestionValidator
from app.core.exceptions import ChapterNotFoundError, InsufficientSourceContentError
from app.services.docx_segmenter import DocxChapterSegmenter


def test_sufficiency_check_rejects_titles_and_metadata_only():
    """Verifies that content with only titles, headings, and metadata fails sufficiency check."""
    metadata_only = """
    Chapter 5: Cell Division
    Page 45
    See Chapter 6 for details
    Total Marks: 100
    Instructions to candidates:
    1. Attempt all questions.
    """
    sanitized = ContentSanitizer.sanitize(metadata_only)
    is_sufficient, words, msg = ContentSanitizer.check_sufficiency(sanitized, "5")
    assert not is_sufficient
    assert words < ContentSanitizer.MIN_SUBSTANTIVE_WORDS

    with pytest.raises(InsufficientSourceContentError):
        ContentSanitizer.validate_or_raise(sanitized, "5")


def test_sufficiency_check_accepts_substantive_educational_content():
    """Verifies that rich educational text passes the sufficiency check."""
    educational_content = """
    Photosynthesis is the biological process by which green plants and certain other organisms transform
    light energy into chemical energy. During photosynthesis in green plants, light energy is captured
    and used to convert water, carbon dioxide, and minerals into oxygen and energy-rich organic compounds.
    Chlorophyll is the green pigment responsible for absorbing light energy in the chloroplasts of plant cells.
    The light-dependent reactions take place in the thylakoid membrane where ATP and NADPH are synthesized.
    The light-independent Calvin cycle occurs in the stroma to produce glucose.
    """
    sanitized = ContentSanitizer.sanitize(educational_content)
    is_sufficient, words, msg = ContentSanitizer.check_sufficiency(sanitized, "Photosynthesis")
    assert is_sufficient
    assert words >= ContentSanitizer.MIN_SUBSTANTIVE_WORDS


def test_question_validator_rejects_metadata_and_structure_questions():
    """Verifies that questions asking about chapter numbers, page numbers, titles, metadata, or exam specs are rejected."""
    source_content = """
    Thermodynamics is the branch of physics that deals with heat, work, and temperature,
    and their relation to energy, radiation, and physical properties of matter.
    The first law of thermodynamics states that the total energy of an isolated system is constant;
    energy can be transformed from one form to another, but can be neither created nor destroyed.
    Entropy is a measure of system disorder.
    """

    candidate_questions = [
        # 1. Chapter number question -> MUST BE REJECTED
        {"text": "What is the chapter number of this topic?", "options": ["A) 5 *", "B) 6", "C) 7", "D) 8"], "markValue": 1},
        # 2. Page number question -> MUST BE REJECTED
        {"text": "On which page is entropy discussed in this textbook?", "options": ["A) Page 45 *", "B) 46", "C) 47", "D) 48"], "markValue": 1},
        # 3. Spelling/title question -> MUST BE REJECTED
        {"text": "What is the spelling of the chapter title?", "options": ["A) Thermodynamics *", "B) Thermodinamics"], "markValue": 1},
        # 4. Exam specification question -> MUST BE REJECTED
        {"text": "How many marks does question 1 carry in this examination paper?", "options": ["A) 1 *", "B) 2", "C) 3", "D) 4"], "markValue": 1},
        # 5. Document metadata question -> MUST BE REJECTED
        {"text": "Who is the author of this textbook?", "options": ["A) Newton *", "B) Einstein", "C) Maxwell", "D) Bohr"], "markValue": 1},
        # 6. Valid educational question -> MUST BE ACCEPTED
        {"text": "State the first law of thermodynamics and explain what it implies about energy conservation.", "markValue": 3},
        # 7. Valid educational MCQ -> MUST BE ACCEPTED
        {"text": "What does entropy measure in a thermodynamic system?", "options": ["A) Temperature", "B) System disorder *", "C) Heat capacity", "D) Mass"], "markValue": 1},
    ]

    accepted, rejected = QuestionValidator.validate_and_filter_questions(candidate_questions, source_content)

    assert len(accepted) == 2
    assert len(rejected) == 5

    # Check that the two accepted are the real physics questions
    accepted_texts = [q["text"] for q in accepted]
    assert any("first law of thermodynamics" in t for t in accepted_texts)
    assert any("entropy" in t for t in accepted_texts)


def test_question_validator_rejects_hallucinated_unsupported_questions():
    """Verifies that questions asking about concepts completely absent from the source content are rejected."""
    source_content = "The nucleus contains protons and neutrons. Electrons orbit around the nucleus in shells."

    unsupported_question = {
        "text": "What is the function of the ribosome in protein synthesis inside the endoplasmic reticulum?",
        "options": ["A) Translation *", "B) Transcription", "C) Replication", "D) Splicing"],
        "markValue": 1
    }

    accepted, rejected = QuestionValidator.validate_and_filter_questions([unsupported_question], source_content)
    assert len(accepted) == 0
    assert len(rejected) == 1
    assert "not grounded" in rejected[0][1]


def test_no_fallback_on_missing_chapter():
    """Verifies that requesting a non-existent chapter raises ChapterNotFoundError with no silent fallback."""
    import io
    import docx

    # Create dummy docx in memory with Chapter 1 only
    doc = docx.Document()
    doc.add_heading("Chapter 1: Foundations of Physics", level=1)
    doc.add_paragraph("Physics is the study of matter and energy.")
    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    # Requesting non-existent Chapter 9 must raise ChapterNotFoundError
    with pytest.raises(ChapterNotFoundError):
        DocxChapterSegmenter.extract_chapter_content(file_bytes, ["Chapter 9"])
